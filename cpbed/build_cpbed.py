# -*- coding: utf-8 -*-
"""Étude de cas - Construction siège social DIGITECH (BATIPROJET). Dossier Word."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG = "/home/user/licence/cpbed/img"
OUT = "/home/user/licence/Etude_de_cas_CPBED_BATIPROJET.docx"

NAVY  = RGBColor(0x1F, 0x4E, 0x79)
NAVY2 = RGBColor(0x2E, 0x6B, 0xA8)
GREY  = RGBColor(0x40, 0x40, 0x40)
LGREY = RGBColor(0x7F, 0x7F, 0x7F)
RED   = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
HDR_BG = "1F4E79"; BAND = "EAF1F8"; WARN_BG = "FDECEA"; OK_BG = "E8F5E9"
GREEN_BG = "C6E0B4"; YEL_BG = "FFE699"; ORA_BG = "F8CBAD"; REDC_BG = "F4827B"

NBSP = " "
def fr(t):
    if t is None: return t
    for a, b in ((" :", NBSP+":"), (" ;", NBSP+";"), (" !", NBSP+"!"),
                 (" ?", NBSP+"?"), (" %", NBSP+"%"), ("« ", "«"+NBSP), (" »", NBSP+"»")):
        t = t.replace(a, b)
    return t

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = GREY
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = 1.13
pf.space_after = Pt(7); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for lvl, sz in [("Heading 1", 15), ("Heading 2", 12.5)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"; st.font.size = Pt(sz); st.font.bold = True
    st.font.color.rgb = NAVY
    st.paragraph_format.space_before = Pt(12 if lvl == "Heading 1" else 8)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True

sec = doc.sections[0]
sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(1.9)
sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

# ---------- helpers ----------
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexc)
    tcPr.append(shd)

def cmargins(cell, t=30, b=30, l=80, r=80):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for tag, v in (("top", t), ("bottom", b), ("start", l), ("end", r)):
        e = OxmlElement("w:" + tag); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)

def ctext(cell, text, bold=False, color=None, size=9.5, align="left"):
    cell.text = ""; p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(fr(str(text))); r.bold = bold; r.font.size = Pt(size); r.font.name = "Calibri"
    if color is not None: r.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def table(rows, headers, widths, aligns=None, fontsize=9.5, band=True,
          caption=None, highlight=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"; t.autofit = False
    for j, h in enumerate(headers):
        ctext(t.rows[0].cells[j], h, bold=True, color=RGBColor(0xFF,0xFF,0xFF),
              size=fontsize, align="center")
        shade(t.rows[0].cells[j], HDR_BG); cmargins(t.rows[0].cells[j])
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        hl = highlight(i, row) if highlight else None
        for j, v in enumerate(row):
            a = aligns[j] if aligns else "left"
            ctext(cells[j], v, size=fontsize, align=a); cmargins(cells[j])
            if hl: shade(cells[j], hl)
            elif band and i % 2 == 1: shade(cells[j], BAND)
    for row in t.rows:
        for j, w in enumerate(widths): row.cells[j].width = Cm(w)
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true"); trPr.append(th)
    if caption: cap(caption)
    return t

def cap(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(9)
    r = p.add_run(fr(text)); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = LGREY

def image(path, w, caption=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(w))
    if caption: cap(caption)

def para(text, bold=False, italic=False, color=None, size=11, align="justify", sa=7, sb=0):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
    r = p.add_run(fr(text)); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    return p

def bullet(text, lead=None):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
    if lead:
        r = p.add_run(fr(lead)); r.bold = True; r.font.size = Pt(11)
    r2 = p.add_run(fr(text)); r2.font.size = Pt(11)
    return p

def h1(t): return doc.add_heading(fr(t), level=1)
def h2(t): return doc.add_heading(fr(t), level=2)

def callout(title, lines, bg=OK_BG, border=GREEN):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0); shade(cell, bg); cmargins(cell, 110, 110, 180, 180); cell.width = Cm(16.6)
    p0 = cell.paragraphs[0]; p0.paragraph_format.space_after = Pt(3)
    r = p0.add_run(fr(title)); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    for ln in lines:
        p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        rr = p.add_run(fr(ln)); rr.font.size = Pt(10.5); rr.font.color.rgb = GREY
    tcPr = cell._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "start", "end"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "18")
        e.set(qn("w:color"), "%02X%02X%02X" % (border[0], border[1], border[2])); bd.append(e)
    tcPr.append(bd)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def hrule(color="1F4E79", size="22", sa=10):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(sa)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), size); bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color); pbdr.append(bot); pPr.append(pbdr)

# =====================================================================
# COUVERTURE (compacte)
# =====================================================================
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ÉTUDE DE CAS — GESTION DE PROJET BÂTIMENT"); r.font.size = Pt(13)
r.font.bold = True; r.font.color.rgb = LGREY
hrule(sa=12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Construction du siège social DIGITECH France")
r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = NAVY
p.paragraph_format.space_after = Pt(4)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Bâtiment tertiaire à énergie positive — Entreprise BATIPROJET")
r.font.size = Pt(14); r.font.italic = True; r.font.color.rgb = NAVY2
p.paragraph_format.space_after = Pt(10)
hrule(sa=14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Dossier de conduite de projet — réponses au travail demandé")
r.font.size = Pt(12); r.font.color.rgb = GREY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Rôle : chef de projet BATIPROJET   ·   17 juin 2026")
r.font.size = Pt(10.5); r.font.italic = True; r.font.color.rgb = LGREY

# Sommaire
doc.add_paragraph().paragraph_format.space_after = Pt(6)
pt = doc.add_paragraph(); pt.paragraph_format.space_after = Pt(4)
r = pt.add_run("Sommaire"); r.font.size = Pt(14); r.bold = True; r.font.color.rgb = NAVY
toc_items = [
    "Question 1 — Note de cadrage",
    "Question 2 — Work Breakdown Structure (WBS)",
    "Question 3 — Planning et diagramme de Gantt",
    "Question 4 — Matrice RACI",
    "Question 5 — Analyse des risques",
    "Question 6 — Analyse réglementaire et normative",
    "Questions de réflexion stratégique",
    "Conclusion — recommandation au comité de direction",
]
def fld(kind, dirty=False):
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
    if dirty: e.set(qn("w:dirty"), "true")
    return e
last = None
for idx, it in enumerate(toc_items):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.2)
    if idx == 0:
        r0 = p.add_run(); instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve"); instr.text = ' TOC \\o "1-1" \\h \\z \\u '
        r0._r.append(fld("begin", True)); r0._r.append(instr); r0._r.append(fld("separate"))
    r = p.add_run(fr(it)); r.font.size = Pt(11); r.font.color.rgb = NAVY
    last = p
last.add_run()._r.append(fld("end"))
para("Sommaire automatique : pour afficher les numéros de page, clic droit dessus puis "
     "« Mettre à jour les champs » (F9).", italic=True, color=LGREY, size=9, sa=2)
doc.add_page_break()

# =====================================================================
# Q1 — NOTE DE CADRAGE
# =====================================================================
h1("Question 1 — Note de cadrage")

h2("Contexte")
para("La société DIGITECH France a retenu l’entreprise générale BATIPROJET pour construire "
     "son nouveau siège social, dans une zone d’activités en pleine expansion. Le programme "
     "porte sur un bâtiment tertiaire moderne de 4 500 m² de bureaux sur 3 niveaux, complété "
     "par des espaces collaboratifs, une salle de conférence de 200 places, un restaurant "
     "d’entreprise, un parking de 120 places avec bornes de recharge, une toiture "
     "photovoltaïque et des espaces verts. Le maître d’ouvrage veut en faire une référence "
     "régionale en matière de performance environnementale (bâtiment à énergie positive).")
para("Le chantier se trouve à proximité immédiate d’un quartier résidentiel dont les riverains "
     "ont déjà fait part de leurs inquiétudes sur les nuisances. À cela s’ajoutent plusieurs "
     "contraintes de marché : hausse du coût des matériaux, tension sur la main-d’œuvre "
     "qualifiée et délais d’approvisionnement longs sur certains équipements techniques.")

h2("Objectifs du projet")
table(
    rows=[
        ["Coût", "Budget maximal de 9 M€ HT"],
        ["Délai", "Livraison sous 18 mois"],
        ["Qualité", "Conformité réglementaire, certification HQE, RE2020, taux de réserves < 2 % à la réception"],
        ["Sécurité", "Zéro accident grave, respect intégral du plan SPS, taux de fréquence inférieur aux moyennes du secteur"],
        ["Environnement", "Réduction des émissions carbone, valorisation des déchets, optimisation des consommations"],
    ],
    headers=["Axe", "Objectif"],
    widths=[3.3, 13.2], aligns=["left", "left"], fontsize=10,
    caption="Objectifs du projet (Document 1).",
)

h2("Livrables attendus")
para("Le projet doit aboutir à la livraison d’un bâtiment complet et opérationnel : les "
     "4 500 m² de bureaux, la salle de conférence, le restaurant, le parking et ses bornes de "
     "recharge, la toiture photovoltaïque et les espaces verts. S’y ajoutent les livrables "
     "documentaires et de conformité : la certification HQE, l’attestation de conformité "
     "RE2020, le dossier des ouvrages exécutés (DOE) et le procès-verbal de réception.")

h2("Parties prenantes")
table(
    rows=[
        ["Maître d’ouvrage (DIGITECH France)", "Finance le projet et valide les décisions"],
        ["Chef de projet (BATIPROJET)", "Pilote le projet des études à la réception"],
        ["Architecte", "Conception architecturale"],
        ["Bureau d’études techniques (BET)", "Études techniques (structure, fluides, thermique)"],
        ["Bureau de contrôle", "Vérification de la conformité"],
        ["Coordinateur SPS", "Sécurité et protection de la santé sur le chantier"],
        ["Conducteur de travaux", "Réalisation et coordination des travaux"],
        ["Sous-traitants", "Exécution des différents lots"],
        ["Mairie", "Instruction du permis et autorisations"],
        ["Futurs utilisateurs / riverains", "Usage du bâtiment / voisinage à concerter"],
    ],
    headers=["Acteur", "Rôle"],
    widths=[6.2, 10.3], aligns=["left", "left"], fontsize=10,
    caption="Principales parties prenantes (Document 3).",
)

h2("Contraintes")
bullet("budget plafonné à 9 M€ HT et délai ferme de 18 mois ;")
bullet("hausse du coût des matériaux de construction ;")
bullet("tension sur les ressources humaines qualifiées ;")
bullet("délais d’approvisionnement importants sur certains équipements techniques ;")
bullet("exigences renforcées en matière de sécurité et d’environnement ;")
bullet("proximité d’un quartier résidentiel (nuisances, attentes des riverains).")

h2("Hypothèses")
bullet("le terrain est disponible, viabilisé et le financement du maître d’ouvrage est sécurisé ;")
bullet("le permis de construire est obtenu dans les délais prévus ;")
bullet("aucun aléa géotechnique majeur n’est rencontré au terrassement ;")
bullet("les entreprises et sous-traitants nécessaires sont disponibles ;")
bullet("les conditions météorologiques restent dans la normale saisonnière.")

h2("Indicateurs de succès")
bullet("livraison dans le délai de 18 mois et coût final ≤ 9 M€ HT ;")
bullet("certification HQE obtenue et conformité RE2020 attestée ;")
bullet("taux de réserves inférieur à 2 % à la réception ;")
bullet("zéro accident grave et taux de fréquence sous la moyenne du secteur ;")
bullet("objectifs de valorisation des déchets et de réduction carbone atteints.")

h2("Facteurs clés de réussite")
bullet("anticiper le permis de construire et les approvisionnements à long délai ;")
bullet("piloter de près le chemin critique et les jalons ;")
bullet("coordonner les lots, idéalement via une maquette BIM ;")
bullet("concerter les riverains et tenir un chantier propre et à faibles nuisances ;")
bullet("sélectionner des entreprises fiables et manager la sécurité au quotidien.")

doc.add_page_break()

# =====================================================================
# Q2 — WBS
# =====================================================================
h1("Question 2 — Work Breakdown Structure (WBS)")
para("Le projet est découpé en six phases, depuis les études jusqu’à la réception, complétées "
     "par un lot transverse de management de projet. Chaque maille correspond à un livrable "
     "identifiable, que l’on peut affecter à un responsable, budgéter et planifier : le "
     "découpage est donc directement exploitable pour le pilotage opérationnel.")
image(IMG + "/wbs.png", 16.6, "WBS du projet — 6 phases et un lot transverse de management.")
para("Le lot transverse (pilotage, qualité HQE/RE2020, sécurité SPS, environnement et BIM) "
     "court sur toute la durée du projet : il ne produit pas un ouvrage mais garantit que les "
     "objectifs de coût, de délai, de qualité, de sécurité et d’environnement sont tenus.")

doc.add_page_break()

# =====================================================================
# Q3 — GANTT
# =====================================================================
h1("Question 3 — Planning et diagramme de Gantt")
para("À partir des durées du Document 2, j’ai défini les liens logiques entre activités. "
     "Trois activités d’études peuvent se chevaucher en amont (la consultation des entreprises "
     "se déroule pendant l’instruction du permis), puis le chantier s’enchaîne de façon "
     "essentiellement séquentielle ; seuls les lots techniques se déroulent en parallèle du "
     "second œuvre.")

h2("Tableau des activités et dépendances")
table(
    rows=[
        ["Études détaillées", "2", "—", "Oui"],
        ["Permis de construire", "3", "Études détaillées", "Oui"],
        ["Consultation des entreprises", "2", "Études détaillées", "Non (marge 1)"],
        ["Installation de chantier", "1", "Permis + Consultation", "Oui"],
        ["Terrassement", "1", "Installation", "Oui"],
        ["Fondations", "2", "Terrassement", "Oui"],
        ["Structure béton", "4", "Fondations", "Oui"],
        ["Couverture / étanchéité", "2", "Structure béton", "Oui"],
        ["Second œuvre", "4", "Couverture / étanchéité", "Oui"],
        ["Lots techniques", "3", "Couverture / étanchéité", "Non (marge 1)"],
        ["Essais et mise en service", "1", "Second œuvre + Lots techniques", "Oui"],
        ["Réception", "1", "Essais et mise en service", "Oui"],
    ],
    headers=["Activité", "Durée (mois)", "Antériorité", "Critique"],
    widths=[5.6, 2.6, 5.6, 2.7], aligns=["left", "center", "left", "center"], fontsize=9.5,
    caption="Durées (Document 2), dépendances retenues et appartenance au chemin critique.",
    highlight=lambda i, r: OK_BG if r[3].startswith("Oui") else None,
)

h2("Diagramme de Gantt")
image(IMG + "/gantt.png", 16.6, "Diagramme de Gantt — chemin critique en rouge, marges en gris.")

h2("Chemin critique et durée")
para("Chemin critique : Études détaillées → Permis de construire → Installation → "
     "Terrassement → Fondations → Structure béton → Couverture/étanchéité → Second œuvre → "
     "Essais → Réception.", bold=True, color=NAVY)
callout("Résultat du planning",
        ["Durée totale du chemin critique = 21 mois.",
         "Or l’objectif de livraison est de 18 mois : le planning prévisionnel dépasse la "
         "cible de 3 mois. Tenir les 18 mois suppose de comprimer le chemin critique "
         "(voir la question de réflexion n° 3)."],
        bg=WARN_BG, border=RED)

h2("Activités susceptibles de générer des retards")
bullet("administratif et hors maîtrise directe de l’entreprise, c’est le premier risque de "
       "dérapage du planning ;", lead="Permis de construire : ")
bullet("4 mois sur le chemin critique, sensible aux intempéries et à la disponibilité des "
       "équipes ;", lead="Structure béton : ")
bullet("4 mois, forte coactivité entre corps d’état et dépendance aux approvisionnements ;",
       lead="Second œuvre : ")
bullet("délais d’approvisionnement longs (panneaux photovoltaïques, CVC, bornes de recharge) "
       "qui peuvent contaminer le planning malgré leur marge.", lead="Lots techniques : ")

doc.add_page_break()

# =====================================================================
# Q4 — RACI
# =====================================================================
h1("Question 4 — Matrice RACI")
para("La matrice ci-dessous répartit les responsabilités sur les huit activités demandées. "
     "Pour chaque activité, un seul acteur est responsable (A), afin que la ligne de décision "
     "reste claire.")
acteurs = ["MOA", "Chef de projet", "Architecte", "BET", "Cond. travaux", "Bureau contrôle", "SPS"]
raci = [
    ["Études de conception",       "A", "C", "R", "R", "I", "C", "I"],
    ["Permis de construire",       "I", "A", "R", "C", "I", "C", "I"],
    ["Consultation des entreprises","C", "A", "C", "R", "I", "I", "I"],
    ["Préparation du chantier",    "I", "A", "I", "C", "R", "C", "C"],
    ["Gros œuvre",                 "I", "A", "I", "C", "R", "C", "C"],
    ["Second œuvre",               "I", "A", "C", "C", "R", "C", "C"],
    ["Essais",                     "I", "A", "I", "C", "R", "C", "I"],
    ["Réception",                  "A", "R", "C", "C", "C", "C", "I"],
]
def raci_hl(i, r):
    return None
t = table(
    rows=raci,
    headers=["Activité"] + acteurs,
    widths=[3.9, 1.85, 2.05, 1.85, 1.35, 1.85, 1.85, 1.0],
    aligns=["left"] + ["center"]*7, fontsize=8.5,
    caption="Matrice RACI des activités du projet.",
)
para("Légende : R = Réalise · A = Responsable (rend des comptes) · C = Consulté · I = Informé.",
     italic=True, size=9.5, color=LGREY)
para("Lecture : le chef de projet est responsable de l’enchaînement opérationnel (du permis "
     "aux essais) ; le maître d’ouvrage reste responsable de la validation de la conception et "
     "de la réception, qui engagent le contrat. Le bureau de contrôle et le coordinateur SPS "
     "sont systématiquement consultés sur les phases travaux.")

doc.add_page_break()

# =====================================================================
# Q5 — RISQUES
# =====================================================================
h1("Question 5 — Analyse des risques")
para("Le registre ci-dessous recense douze risques couvrant les coûts, les délais, la qualité, "
     "la sécurité, l’environnement, le réglementaire, les ressources humaines, les "
     "sous-traitants et les approvisionnements. La probabilité (P) et l’impact (I) sont notés "
     "de 1 à 4 ; la criticité est leur produit (P × I).")
risks = [
    ["R1", "Dépassement du délai de 18 mois (chemin critique à 21 mois)", "4", "4", "16",
     "Comprimer le chemin critique, anticiper permis et appros, préfabrication, pilotage serré"],
    ["R2", "Hausse du coût des matériaux", "4", "3", "12",
     "Clauses de révision de prix, achats anticipés, marchés à prix ferme"],
    ["R4", "Délais d’approvisionnement (PV, CVC, bornes IRVE)", "3", "4", "12",
     "Commandes anticipées, fournisseurs alternatifs, stock tampon"],
    ["R3", "Pénurie de main-d’œuvre qualifiée", "3", "3", "9",
     "Sécuriser les contrats de sous-traitance, planifier les ressources"],
    ["R6", "Nuisances et recours des riverains", "3", "3", "9",
     "Concertation, charte chantier propre, horaires encadrés, communication"],
    ["R9", "Défaillance d’un sous-traitant", "3", "3", "9",
     "Sélection rigoureuse, cautions, clauses contractuelles, suivi rapproché"],
    ["R5", "Refus ou retard du permis de construire", "2", "4", "8",
     "Dossier solide, pré-consultation mairie/ABF, marge de sécurité"],
    ["R7", "Non-obtention HQE / non-conformité RE2020", "2", "4", "8",
     "AMO HQE, BET thermique, contrôles dès la phase études"],
    ["R8", "Accident grave sur le chantier", "2", "4", "8",
     "Plan SPS strict, formations, contrôles et présence SPS"],
    ["R10", "Intempéries retardant le gros œuvre", "3", "2", "6",
     "Planning saisonnier, intempéries provisionnées au planning"],
    ["R11", "Taux de réserves > 2 % à la réception", "2", "3", "6",
     "Contrôles intermédiaires, opérations préalables à la réception anticipées"],
    ["R12", "Mauvaise gestion des déchets (loi AGEC)", "2", "3", "6",
     "SOGED, tri sur site, valorisation et traçabilité des déchets"],
]
def risk_hl(i, r):
    c = int(r[4])
    if c >= 13: return REDC_BG
    if c >= 9: return ORA_BG
    if c >= 5: return YEL_BG
    return GREEN_BG
table(
    rows=risks,
    headers=["#", "Risque", "P", "I", "Crit.", "Action de traitement"],
    widths=[0.9, 5.3, 0.8, 0.8, 1.1, 7.6],
    aligns=["center", "left", "center", "center", "center", "left"], fontsize=8.7,
    band=False,
    caption="Registre des risques (criticité colorée : vert faible, jaune modérée, orange élevée, rouge majeure).",
    highlight=risk_hl,
)
image(IMG + "/risques.png", 12.5, "Matrice des risques (P × I).")
callout("Les trois risques majeurs",
        ["1. Le dépassement du délai (R1) — le planning prévisionnel est à 21 mois pour un "
         "objectif de 18.",
         "2. La hausse du coût des matériaux (R2) — elle menace directement le budget de 9 M€.",
         "3. Les délais d’approvisionnement des équipements techniques (R4) — sur des lots à "
         "long délai (photovoltaïque, CVC, bornes)."],
        bg=BAND, border=NAVY2)

doc.add_page_break()

# =====================================================================
# Q6 — RÉGLEMENTAIRE
# =====================================================================
h1("Question 6 — Analyse réglementaire et normative")
para("Le projet est soumis à un ensemble d’exigences réglementaires et normatives qu’il faut "
     "intégrer dès la conception. Un point mérite d’être souligné : avec une salle de "
     "conférence de 200 places et un restaurant d’entreprise, le bâtiment relève de la "
     "réglementation des établissements recevant du public (ERP), ce qui renforce les "
     "exigences en matière de sécurité incendie et d’accessibilité.")
table(
    rows=[
        ["Urbanisme", "Code de l’urbanisme, PLU, permis de construire",
         "Implantation, hauteurs, stationnement, dépôt et instruction en mairie"],
        ["Construction", "RE2020, Eurocodes, DTU",
         "Performance énergétique et carbone, dimensionnement de la structure, règles de l’art"],
        ["Sécurité", "Code du travail, coordination SPS, PGC/PPSPS",
         "Prévention des risques, plan SPS, objectif zéro accident grave"],
        ["Accessibilité", "Réglementation PMR, obligations ERP",
         "Accès et circulations adaptés ; salle de conférence et restaurant = ERP"],
        ["Incendie", "Règlement ERP, SSI, désenfumage, compartimentage",
         "Système de sécurité incendie, dégagements et issues de secours"],
        ["Environnement", "Loi AGEC, gestion des déchets, bilan carbone",
         "Tri et valorisation des déchets, économie circulaire, mesure du carbone"],
        ["Normes / référentiels", "HQE, ISO 9001, ISO 14001, ISO 45001, ISO 19650 (BIM)",
         "Management qualité, environnement et sécurité ; maquette numérique BIM"],
    ],
    headers=["Domaine", "Textes et référentiels", "Application au projet"],
    widths=[3.1, 5.5, 7.9], aligns=["left", "left", "left"], fontsize=9.3,
    caption="Principales réglementations et normes applicables.",
)
para("Le respect de ces exigences conditionne à la fois l’autorisation de construire "
     "(permis, ERP), la conformité de l’ouvrage (RE2020, Eurocodes, DTU) et l’atteinte des "
     "objectifs environnementaux affichés par le maître d’ouvrage (HQE, loi AGEC).")

doc.add_page_break()

# =====================================================================
# QUESTIONS DE RÉFLEXION
# =====================================================================
h1("Questions de réflexion stratégique")

h2("1. Les trois risques majeurs du projet")
para("Le délai (le planning prévisionnel ressort à 21 mois pour un objectif de 18), la hausse "
     "du coût des matériaux qui pèse directement sur le budget de 9 M€, et les délais "
     "d’approvisionnement des équipements techniques (photovoltaïque, CVC, bornes de recharge).")

h2("2. Impact d’une hausse de 20 % du coût des matériaux")
para("Sur une opération de ce type, les matériaux représentent de l’ordre de la moitié du coût "
     "de construction. Une hausse de 20 % sur cette part se traduit donc par un surcoût "
     "d’environ 10 % du budget total, soit près de 0,9 M€ sur 9 M€ : l’enveloppe plafond serait "
     "dépassée. Les leviers sont les clauses de révision de prix, l’achat anticipé des "
     "matériaux sensibles et l’optimisation de la conception (quantités, choix techniques).")

h2("3. Sécuriser le respect du délai de 18 mois")
para("Puisque le chemin critique ressort à 21 mois, il faut le comprimer de 3 mois. "
     "Concrètement : déposer le permis au plus tôt et instruire les consultations en parallèle, "
     "recourir à la préfabrication pour la structure, faire chevaucher le démarrage des lots "
     "techniques et du second œuvre, commander dès le départ les équipements à long délai, "
     "provisionner les intempéries et piloter les jalons critiques en réunion de chantier "
     "hebdomadaire. Le BIM aide à fiabiliser cette coordination.")

h2("4. Indicateurs de pilotage pour le comité de pilotage")
para("Avancement physique par rapport au planning (jalons tenus), indices de performance "
     "délai et coût issus de la valeur acquise (SPI et CPI), coût engagé par rapport au budget, "
     "taux de réserves, indicateurs de sécurité (taux de fréquence et de gravité, "
     "presqu’accidents), indicateurs environnementaux (taux de valorisation des déchets, bilan "
     "carbone) et suivi du registre des risques.")

h2("5. Apports du BIM dans ce projet")
para("La maquette numérique partagée (ISO 19650) améliore la coordination entre lots, permet "
     "de détecter les conflits avant le chantier, fiabilise les quantitatifs et donc le budget, "
     "et facilite l’exploitation future grâce à un DOE numérique. Au final, moins d’erreurs et "
     "de reprises, donc des gains de délai, de coût et de qualité.")

h2("6. Concilier performance économique, environnementale et sécurité")
para("La clé est de raisonner en coût global plutôt qu’en coût d’achat : un bâtiment HQE/RE2020 "
     "coûte un peu plus cher à la construction mais moins cher à l’usage. Il faut intégrer ces "
     "trois dimensions dès les études (conception bioclimatique, matériaux bas carbone "
     "optimisés, sécurité prise en compte dans le planning via le SPS), et arbitrer les "
     "compromis à l’aide du BIM. Bien anticipée, la sécurité ne coûte pas de temps : ce sont "
     "les accidents qui en font perdre.")

doc.add_page_break()

# =====================================================================
# CONCLUSION
# =====================================================================
h1("Conclusion — recommandation au comité de direction")
para("Le projet est techniquement réalisable et son cadrage est solide. Deux réserves doivent "
     "toutefois être levées avant le lancement définitif. D’abord le délai : le planning "
     "prévisionnel ressort à 21 mois, soit 3 de plus que l’objectif ; il faut valider un plan "
     "de compression (anticipation du permis et des approvisionnements, préfabrication, "
     "chevauchement des lots) avant de s’engager sur les 18 mois. Ensuite le budget, exposé à "
     "la hausse des matériaux, qu’il faut sécuriser par des achats anticipés et des clauses de "
     "révision.")
para("Recommandation : lancer le projet sous conditions — plan de compression du délai validé, "
     "marché des matériaux sécurisé et registre des risques suivi en comité de pilotage. À ces "
     "conditions, les objectifs de coût, de délai, de qualité environnementale et de sécurité "
     "peuvent être tenus.", bold=True, color=NAVY)

# =====================================================================
# Pied de page (numéro de page corrigé : cache ENTRE separate et end)
# =====================================================================
def add_field(paragraph, code):
    r1 = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = code
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    r1._r.append(b); r1._r.append(i); r1._r.append(s)
    rc = paragraph.add_run("1"); rc.font.size = Pt(8.5); rc.font.color.rgb = LGREY
    r3 = paragraph.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    r3._r.append(e)

section = doc.sections[0]
section.different_first_page_header_footer = True
ftr = section.footer; ftr.is_linked_to_previous = False
fp = ftr.paragraphs[0]; fp.text = ""
tabs = fp.paragraph_format.tab_stops
tabs.add_tab_stop(Cm(8.1), WD_TAB_ALIGNMENT.CENTER)
tabs.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT)
r = fp.add_run("Étude de cas — Siège social DIGITECH"); r.font.size = Pt(8.5); r.font.color.rgb = LGREY
fp.add_run("\t"); rmid = fp.add_run("Page "); rmid.font.size = Pt(8.5); rmid.font.color.rgb = LGREY
add_field(fp, " PAGE "); rs = fp.add_run(" / "); rs.font.size = Pt(8.5); rs.font.color.rgb = LGREY
add_field(fp, " NUMPAGES "); fp.add_run("\t")
rr = fp.add_run("BATIPROJET"); rr.font.size = Pt(8.5); rr.font.color.rgb = LGREY

# mise à jour des champs (sommaire) à l'ouverture
settings = doc.settings.element
upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); settings.append(upd)

doc.save(OUT)
print("Dossier généré :", OUT)
