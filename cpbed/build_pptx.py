# -*- coding: utf-8 -*-
"""Diaporama de soutenance - Étude de cas BATIPROJET / DIGITECH (bâtiment)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

IMG = "/home/user/licence/cpbed/img"
OUT = "/home/user/licence/Etude_de_cas_CPBED_Presentation.pptx"

NAVY   = RGBColor(0x1F, 0x4E, 0x79)
NAVY2  = RGBColor(0x2E, 0x6B, 0xA8)
LBLUE  = RGBColor(0xEA, 0xF1, 0xF8)
GREY   = RGBColor(0x40, 0x40, 0x40)
LGREY  = RGBColor(0x8A, 0x8A, 0x8A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
PAGE = [1]

NBSP = " "
def fr(t):
    for a, b in ((" :", NBSP+":"), (" ;", NBSP+";"), (" !", NBSP+"!"),
                 (" ?", NBSP+"?"), (" %", NBSP+"%"), ("« ", "«"+NBSP), (" »", NBSP+"»")):
        t = t.replace(a, b)
    return t

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    return s

def rect(s, l, t, w, h, color):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def rrect(s, l, t, w, h, color):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color; sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def txt(s, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None: p.space_after = Pt(ln["space_after"])
        p.space_before = Pt(ln.get("space_before", 0))
        if ln.get("level"): p.level = ln["level"]
        r = p.add_run(); r.text = fr(ln["text"]); f = r.font
        f.size = Pt(ln.get("size", 18)); f.bold = ln.get("bold", False)
        f.italic = ln.get("italic", False); f.name = "Calibri"
        f.color.rgb = ln.get("color", GREY)
    return tb

def title_bar(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.25), NAVY)
    rect(s, 0, Inches(1.25), SW, Pt(3), ORANGE)
    txt(s, Inches(0.55), Inches(0.18), Inches(11.8), Inches(0.35),
        [{"text": kicker.upper(), "size": 12, "bold": True, "color": RGBColor(0xBE,0xD3,0xE9)}])
    txt(s, Inches(0.55), Inches(0.46), Inches(12.2), Inches(0.7),
        [{"text": title, "size": 25, "bold": True, "color": WHITE}])

def footer(s):
    txt(s, Inches(0.55), Inches(7.05), Inches(7), Inches(0.3),
        [{"text": "Étude de cas — Siège social DIGITECH (BATIPROJET)", "size": 9, "color": LGREY}])
    txt(s, Inches(11.3), Inches(7.05), Inches(1.5), Inches(0.3),
        [{"text": str(PAGE[0]), "size": 9, "color": LGREY, "align": PP_ALIGN.RIGHT}])
    PAGE[0] += 1

def pic_centered(s, path, top, max_w, max_h, cx=None):
    im = Image.open(path); iw, ih = im.size
    w = max_w; h = Emu(int(w * ih / iw))
    if h > max_h:
        h = max_h; w = Emu(int(h * iw / ih))
    left = Emu(int((SW - w) / 2)) if cx is None else cx
    s.shapes.add_picture(path, left, top, width=w, height=h)
    return w, h

def kpi(s, l, t, w, value, label, color):
    rrect(s, l, t, w, Inches(1.4), LBLUE)
    txt(s, l, t+Inches(0.18), w, Inches(0.7),
        [{"text": value, "size": 26, "bold": True, "color": color, "align": PP_ALIGN.CENTER}])
    txt(s, l, t+Inches(0.9), w, Inches(0.45),
        [{"text": label, "size": 12, "color": GREY, "align": PP_ALIGN.CENTER}])

def bullets(s, l, t, w, h, items, size=18, gap=10):
    lines = []
    for it in items:
        text, lvl = it if isinstance(it, tuple) else (it, 0)
        lines.append({"text": ("•  " if lvl == 0 else "–  ") + text,
                      "size": size if lvl == 0 else size-2,
                      "color": GREY if lvl == 0 else LGREY,
                      "space_after": gap, "level": lvl})
    txt(s, l, t, w, h, lines)

# =====================================================================
# 1. TITRE
# =====================================================================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(3.05), SW, Pt(3), ORANGE)
rect(s, 0, 0, Inches(0.22), SH, ORANGE)
txt(s, Inches(1.0), Inches(0.9), Inches(11), Inches(0.5),
    [{"text": "ÉTUDE DE CAS — GESTION DE PROJET BÂTIMENT", "size": 15, "bold": True,
      "color": RGBColor(0xBE,0xD3,0xE9)}])
txt(s, Inches(1.0), Inches(1.5), Inches(11.5), Inches(1.7),
    [{"text": "Siège social DIGITECH", "size": 50, "bold": True, "color": WHITE}])
txt(s, Inches(1.0), Inches(3.3), Inches(11.5), Inches(1.3),
    [{"text": "Bâtiment tertiaire à énergie positive", "size": 24, "color": RGBColor(0xDD,0xE8,0xF3)},
     {"text": "Entreprise générale BATIPROJET", "size": 18, "italic": True,
      "color": RGBColor(0xBE,0xD3,0xE9), "space_before": 6}])
txt(s, Inches(1.0), Inches(6.1), Inches(11.5), Inches(0.6),
    [{"text": "Chef de projet BATIPROJET — 17 juin 2026", "size": 13, "color": RGBColor(0x9F,0xBA,0xD6)}])

# =====================================================================
# 2. PLAN
# =====================================================================
s = slide(); title_bar(s, "Déroulé", "Plan de la présentation"); footer(s)
plan = [
    ("1", "Contexte et cadrage", "L’opération, les objectifs, les contraintes"),
    ("2", "Découpage du projet (WBS)", "Les phases et les livrables"),
    ("3", "Planning et Gantt", "Dépendances, chemin critique, délai"),
    ("4", "Responsabilités (RACI)", "Qui fait quoi"),
    ("5", "Risques", "Registre et risques majeurs"),
    ("6", "Réglementaire et recommandation", "Conformité et décision de lancement"),
]
y = Inches(1.75)
for num, titre, sub in plan:
    rrect(s, Inches(0.9), y, Inches(0.75), Inches(0.75), NAVY)
    txt(s, Inches(0.9), y+Inches(0.12), Inches(0.75), Inches(0.5),
        [{"text": num, "size": 22, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    txt(s, Inches(1.9), y+Inches(0.02), Inches(10.5), Inches(0.4),
        [{"text": titre, "size": 19, "bold": True, "color": NAVY}])
    txt(s, Inches(1.9), y+Inches(0.42), Inches(10.5), Inches(0.35),
        [{"text": sub, "size": 13, "color": LGREY}])
    y = y + Inches(0.83)

# =====================================================================
# 3. CONTEXTE
# =====================================================================
s = slide(); title_bar(s, "1 · Cadrage", "Contexte et enjeu"); footer(s)
bullets(s, Inches(0.6), Inches(1.7), Inches(7.1), Inches(5),
    [ "DIGITECH France confie à BATIPROJET son nouveau siège social.",
      ("Bâtiment tertiaire BEPOS : 4 500 m² de bureaux sur 3 niveaux, salle de "
       "conférence de 200 places, restaurant, parking de 120 places, toiture "
       "photovoltaïque.", 1),
      "Objectif du maître d’ouvrage : une référence régionale en performance environnementale.",
      "Chantier proche d’un quartier résidentiel : riverains à concerter.",
    ], size=16.5, gap=12)
rrect(s, Inches(8.0), Inches(1.85), Inches(4.7), Inches(2.0), LBLUE)
txt(s, Inches(8.0), Inches(2.0), Inches(4.7), Inches(0.4),
    [{"text": "LE CADRE", "size": 13, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER}])
txt(s, Inches(8.2), Inches(2.5), Inches(4.3), Inches(1.3),
    [{"text": "9 M€ HT — budget plafond", "size": 18, "bold": True, "color": GREY, "space_after": 8},
     {"text": "18 mois — délai de livraison", "size": 18, "bold": True, "color": GREY}])
rrect(s, Inches(8.0), Inches(4.05), Inches(4.7), Inches(2.3), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(8.2), Inches(4.25), Inches(4.3), Inches(2.0),
    [{"text": "Contexte tendu", "size": 14, "bold": True, "color": RED, "space_after": 6},
     {"text": "Hausse du coût des matériaux, tension sur la main-d’œuvre, délais "
      "d’approvisionnement longs.", "size": 14, "color": GREY}])

# =====================================================================
# 4. OBJECTIFS & CONTRAINTES
# =====================================================================
s = slide(); title_bar(s, "1 · Cadrage", "Objectifs et contraintes"); footer(s)
kpi(s, Inches(0.6), Inches(1.65), Inches(3.9), "9 M€ HT", "Budget maximal", NAVY)
kpi(s, Inches(4.7), Inches(1.65), Inches(3.9), "18 mois", "Délai de livraison", NAVY)
kpi(s, Inches(8.8), Inches(1.65), Inches(3.9), "HQE · RE2020", "Qualité environnementale", NAVY)
txt(s, Inches(0.6), Inches(3.45), Inches(6), Inches(0.5),
    [{"text": "Aussi visés", "size": 15, "bold": True, "color": NAVY}])
bullets(s, Inches(0.6), Inches(3.95), Inches(6.1), Inches(3),
    [ "Sécurité : zéro accident grave, plan SPS respecté.",
      "Environnement : carbone réduit, déchets valorisés.",
      "Qualité : taux de réserves < 2 % à la réception.",
    ], size=15, gap=9)
txt(s, Inches(7.0), Inches(3.45), Inches(6), Inches(0.5),
    [{"text": "Principales contraintes", "size": 15, "bold": True, "color": NAVY}])
bullets(s, Inches(7.0), Inches(3.95), Inches(6.1), Inches(3),
    [ "Hausse du coût des matériaux.",
      "Tension sur la main-d’œuvre qualifiée.",
      "Délais d’approvisionnement des équipements techniques.",
      "Proximité résidentielle (nuisances).",
    ], size=15, gap=9)

# =====================================================================
# 5. WBS
# =====================================================================
s = slide(); title_bar(s, "2 · WBS", "Découpage du projet"); footer(s)
txt(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.6),
    [{"text": "Six phases (des études à la réception) et un lot transverse de management. "
      "Chaque maille = un livrable pilotable.", "size": 15, "color": GREY}])
pic_centered(s, IMG + "/wbs.png", Inches(2.2), Inches(12.4), Inches(4.7))

# =====================================================================
# 6. PLANNING & GANTT
# =====================================================================
s = slide(); title_bar(s, "3 · Planning", "Planning, Gantt et chemin critique"); footer(s)
pic_centered(s, IMG + "/gantt.png", Inches(1.55), Inches(8.7), Inches(4.6), cx=Inches(0.4))
rrect(s, Inches(9.35), Inches(1.55), Inches(3.5), Inches(2.5), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(9.5), Inches(1.72), Inches(3.2), Inches(2.3),
    [{"text": "DÉLAI", "size": 13, "bold": True, "color": RED, "align": PP_ALIGN.CENTER, "space_after": 4},
     {"text": "21 mois", "size": 30, "bold": True, "color": RED, "align": PP_ALIGN.CENTER, "space_after": 2},
     {"text": "pour un objectif de 18 mois", "size": 13, "color": GREY, "align": PP_ALIGN.CENTER, "space_after": 6},
     {"text": "→ à comprimer de 3 mois", "size": 13, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER}])
txt(s, Inches(9.35), Inches(4.3), Inches(3.5), Inches(0.4),
    [{"text": "Activités à risque de retard", "size": 13.5, "bold": True, "color": NAVY}])
bullets(s, Inches(9.35), Inches(4.8), Inches(3.6), Inches(2.3),
    [ "Permis de construire", "Structure béton (4 mois)",
      "Second œuvre (coactivité)", "Lots techniques (appros)"],
    size=12.5, gap=6)

# =====================================================================
# 7. RACI
# =====================================================================
s = slide(); title_bar(s, "4 · Responsabilités", "Matrice RACI"); footer(s)
acteurs = ["MOA", "Chef\nprojet", "Archi.", "BET", "Cond.\ntrav.", "Bureau\ncontr.", "SPS"]
raci = [
    ["Études de conception",        "A", "C", "R", "R", "I", "C", "I"],
    ["Permis de construire",        "I", "A", "R", "C", "I", "C", "I"],
    ["Consultation entreprises",    "C", "A", "C", "R", "I", "I", "I"],
    ["Préparation du chantier",     "I", "A", "I", "C", "R", "C", "C"],
    ["Gros œuvre",                  "I", "A", "I", "C", "R", "C", "C"],
    ["Second œuvre",                "I", "A", "C", "C", "R", "C", "C"],
    ["Essais",                      "I", "A", "I", "C", "R", "C", "I"],
    ["Réception",                   "A", "R", "C", "C", "C", "C", "I"],
]
nrows = len(raci) + 1; ncols = 8
tbl_l, tbl_t = Inches(0.6), Inches(1.6)
tbl_w, tbl_h = Inches(12.1), Inches(4.55)
gtab = s.shapes.add_table(nrows, ncols, tbl_l, tbl_t, tbl_w, tbl_h).table
gtab.columns[0].width = Inches(3.5)
for j in range(1, ncols):
    gtab.columns[j].width = Inches((12.1 - 3.5) / 7)
# header
hdr = ["Activité"] + acteurs
for j, htxt in enumerate(hdr):
    c = gtab.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
    c.margin_top = Pt(1); c.margin_bottom = Pt(1)
    tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = htxt; r.font.size = Pt(11); r.font.bold = True
    r.font.color.rgb = WHITE; r.font.name = "Calibri"
COL = {"R": RGBColor(0x2E,0x7D,0x32), "A": RGBColor(0xC0,0x39,0x2B),
       "C": RGBColor(0x8A,0x8A,0x8A), "I": RGBColor(0xB0,0xB0,0xB0)}
for i, row in enumerate(raci, start=1):
    for j, val in enumerate(row):
        c = gtab.cell(i, j); c.fill.solid()
        c.fill.fore_color.rgb = LBLUE if (i % 2 == 0) else WHITE
        c.margin_top = Pt(1); c.margin_bottom = Pt(1); c.margin_left = Pt(4)
        tf = c.text_frame; p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        r = p.add_run(); r.text = val; r.font.name = "Calibri"
        if j == 0:
            r.font.size = Pt(11); r.font.color.rgb = GREY
        else:
            r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = COL.get(val, GREY)
txt(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.5),
    [{"text": "R = Réalise   ·   A = Responsable (rend des comptes)   ·   C = Consulté   ·   I = Informé",
      "size": 12, "italic": True, "color": LGREY}])

# =====================================================================
# 8. RISQUES
# =====================================================================
s = slide(); title_bar(s, "5 · Risques", "Analyse des risques"); footer(s)
im = Image.open(IMG + "/risques.png"); iw, ih = im.size
h = Inches(5.0); w = Emu(int(h * iw / ih))
s.shapes.add_picture(IMG + "/risques.png", Inches(0.5), Inches(1.7), width=w, height=h)
txt(s, Inches(7.9), Inches(1.7), Inches(5.0), Inches(0.5),
    [{"text": "3 risques majeurs", "size": 16, "bold": True, "color": NAVY}])
prio = [("R1", "Dépassement du délai (21 mois vs 18)", "16", RED),
        ("R2", "Hausse du coût des matériaux", "12", ORANGE),
        ("R4", "Délais d’approvisionnement (PV, CVC)", "12", ORANGE)]
y = Inches(2.35)
for code, lab, c, col in prio:
    rrect(s, Inches(7.9), y, Inches(4.95), Inches(1.05), LBLUE)
    rect(s, Inches(7.9), y, Inches(0.16), Inches(1.05), col)
    txt(s, Inches(8.15), y+Inches(0.1), Inches(3.5), Inches(0.9),
        [{"text": code, "size": 14, "bold": True, "color": col, "space_after": 2},
         {"text": lab, "size": 12.5, "color": GREY}])
    txt(s, Inches(11.85), y+Inches(0.22), Inches(0.9), Inches(0.6),
        [{"text": c, "size": 22, "bold": True, "color": col, "align": PP_ALIGN.CENTER}])
    y = y + Inches(1.2)
txt(s, Inches(7.9), Inches(6.0), Inches(5.0), Inches(0.5),
    [{"text": "Registre complet : 12 risques cotés P × I dans le dossier.",
      "size": 12, "italic": True, "color": LGREY}])

# =====================================================================
# 9. RÉGLEMENTAIRE
# =====================================================================
s = slide(); title_bar(s, "6 · Conformité", "Analyse réglementaire et normative"); footer(s)
reg = [
    ("Urbanisme", "Permis de construire · PLU · Code de l’urbanisme"),
    ("Construction", "RE2020 · Eurocodes · DTU"),
    ("Sécurité", "Code du travail · coordination SPS"),
    ("Accessibilité", "PMR · obligations ERP"),
    ("Incendie", "SSI · désenfumage · compartimentage"),
    ("Environnement", "Loi AGEC · déchets · bilan carbone"),
    ("Normes", "HQE · ISO 9001 / 14001 / 45001 · ISO 19650 (BIM)"),
]
y = Inches(1.7)
for dom, txts in reg:
    rrect(s, Inches(0.6), y, Inches(3.2), Inches(0.62), NAVY)
    txt(s, Inches(0.6), y+Inches(0.12), Inches(3.2), Inches(0.4),
        [{"text": dom, "size": 13.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    txt(s, Inches(4.0), y+Inches(0.13), Inches(8.8), Inches(0.4),
        [{"text": txts, "size": 14, "color": GREY}])
    y = y + Inches(0.7)
rrect(s, Inches(0.6), Inches(6.65), Inches(12.2), Inches(0.62), RGBColor(0xFF,0xF4,0xE0))
txt(s, Inches(0.8), Inches(6.78), Inches(12), Inches(0.4),
    [{"text": "Point clé : salle de conférence (200 places) + restaurant = ERP → règles "
      "incendie et accessibilité renforcées.", "size": 12.5, "bold": True, "color": RGBColor(0x9C,0x6A,0x00)}])

# =====================================================================
# 10. RÉFLEXION STRATÉGIQUE
# =====================================================================
s = slide(); title_bar(s, "Réflexion", "Points clés stratégiques"); footer(s)
qa = [
    ("Hausse de +20 % des matériaux ?",
     "≈ +0,9 M€ sur 9 M€ → l’enveloppe est dépassée. Parades : achats anticipés, clauses de révision."),
    ("Sécuriser le délai de 18 mois ?",
     "Comprimer le chemin critique : permis anticipé, préfabrication, lots techniques en parallèle, appros tôt."),
    ("Indicateurs pour le COPIL ?",
     "Avancement vs planning, SPI/CPI (valeur acquise), coût engagé, sécurité (TF/TG), déchets valorisés."),
    ("Apport du BIM ?",
     "Coordination des lots, détection des conflits, quantitatifs fiables, DOE numérique → gains délai/coût/qualité."),
]
y = Inches(1.7)
for q, a in qa:
    txt(s, Inches(0.7), y, Inches(12), Inches(0.4),
        [{"text": q, "size": 16, "bold": True, "color": NAVY}])
    txt(s, Inches(0.9), y+Inches(0.42), Inches(11.8), Inches(0.6),
        [{"text": a, "size": 14, "color": GREY}])
    y = y + Inches(1.25)

# =====================================================================
# 11. CONCLUSION
# =====================================================================
s = slide(); title_bar(s, "Synthèse", "Recommandation au comité de direction"); footer(s)
txt(s, Inches(0.6), Inches(1.65), Inches(12), Inches(0.6),
    [{"text": "Le projet est solide et réalisable, mais deux réserves doivent être levées avant "
      "le lancement.", "size": 18, "bold": True, "color": NAVY}])
rrect(s, Inches(0.6), Inches(2.6), Inches(6.0), Inches(2.4), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(0.8), Inches(2.8), Inches(5.6), Inches(2.1),
    [{"text": "Délai", "size": 15, "bold": True, "color": RED, "space_after": 4},
     {"text": "Planning à 21 mois pour un objectif de 18. Valider un plan de compression "
      "(permis, préfabrication, chevauchement des lots).", "size": 14, "color": GREY}])
rrect(s, Inches(6.9), Inches(2.6), Inches(6.0), Inches(2.4), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(7.1), Inches(2.8), Inches(5.6), Inches(2.1),
    [{"text": "Budget", "size": 15, "bold": True, "color": RED, "space_after": 4},
     {"text": "Exposé à la hausse des matériaux. Sécuriser par des achats anticipés et des "
      "clauses de révision de prix.", "size": 14, "color": GREY}])
rrect(s, Inches(0.6), Inches(5.25), Inches(12.3), Inches(1.2), RGBColor(0xE8,0xF5,0xE9))
txt(s, Inches(0.8), Inches(5.45), Inches(11.9), Inches(0.9),
    [{"text": "Recommandation : lancer sous conditions — plan de compression du délai validé, "
      "marché des matériaux sécurisé et registre des risques suivi en comité de pilotage.",
      "size": 15, "bold": True, "color": GREEN}])

prs.save(OUT)
print("Présentation générée :", OUT, "—", len(prs.slides._sldIdLst), "diapos")
