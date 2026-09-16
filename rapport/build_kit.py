# -*- coding: utf-8 -*-
"""Assemble les fichiers Markdown du rapport en un document Word mis en forme."""
import glob, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = [os.path.join(os.path.dirname(__file__), 'kit_soutenance.md')]
OUT = os.path.join(os.path.dirname(__file__), 'Kit_soutenance_et_pieces_manquantes.docx')

BLUE = RGBColor(0x00, 0x3A, 0x70)
GREY = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# --- Mise en page ---
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(1.8); s.right_margin = Cm(1.6)

st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
pf = st.paragraph_format
pf.space_after = Pt(3); pf.space_before = Pt(0); pf.line_spacing = 1.0
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, color, before, after in (
        ('Heading 1', 15, BLUE, 10, 5),
        ('Heading 2', 12.5, BLUE, 9, 4),
        ('Heading 3', 11, BLUE, 7, 3),
        ('Heading 4', 10, GREY, 6, 2)):
    s = doc.styles[name]
    s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = True
    s.font.color.rgb = color
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

# --- Pied de page avec numéro ---
def add_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ''
    run = p.add_run()
    for el, attr in (('w:fldChar', {'w:fldCharType': 'begin'}),
                     ('w:instrText', None),
                     ('w:fldChar', {'w:fldCharType': 'end'})):
        e = OxmlElement(el)
        if el == 'w:instrText':
            e.set(qn('xml:space'), 'preserve'); e.text = ' PAGE '
        else:
            for k, v in attr.items():
                e.set(qn(k), v)
        run._r.append(e)
    p.runs[0].font.size = Pt(9)
add_page_numbers(doc.sections[0])

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|<sub>.+?</sub>|<sup>.+?</sup>)')

def add_runs(par, text):
    text = text.replace('\\_', '_')
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            r = par.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        elif part.startswith('<sub>'):
            r = par.add_run(part[5:-6]); r.font.subscript = True
        elif part.startswith('<sup>'):
            r = par.add_run(part[5:-6]); r.font.superscript = True
        else:
            par.add_run(part)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def split_row(line):
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]

def tight_cell_margins(t):
    tblPr = t._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, w in (('top', 10), ('left', 60), ('bottom', 10), ('right', 60)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), str(w)); e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)

def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    tight_cell_margins(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_runs(p, h)
        for r in p.runs:
            r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, '003A70')
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for i, val in enumerate(row[:len(header)]):
            c = cells[i]; c.text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, val)
            for r in p.runs:
                r.font.size = Pt(8)
        if ri % 2 == 1:
            for c in cells:
                shade(c, 'EEF2F7')
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.space_before = Pt(0)
    for r0 in sp.runs:
        r0.font.size = Pt(4)

def add_code(lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.4); pf.space_before = Pt(4); pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('\n'.join(lines))
    r.font.name = 'Consolas'; r.font.size = Pt(7.5)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F4F5F7')
    pPr.append(shd)

# --- Lecture et assemblage ---
text = []
for f in SRC:
    text.append(open(f, encoding='utf-8').read())
raw = '\n\n'.join(text).split('\n')

i = 0
first_heading = True
while i < len(raw):
    line = raw[i].rstrip()
    stripped = line.strip()

    if stripped == 'PAGEBREAK':
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        i += 1; continue

    if stripped.startswith('```'):
        block = []; i += 1
        while i < len(raw) and not raw[i].strip().startswith('```'):
            block.append(raw[i]); i += 1
        i += 1
        add_code(block); continue

    if stripped.startswith('|') and i + 1 < len(raw) and re.match(r'^\|[\s:\-|]+\|$', raw[i+1].strip()):
        rows = [split_row(stripped)]
        i += 2
        while i < len(raw) and raw[i].strip().startswith('|'):
            rows.append(split_row(raw[i])); i += 1
        add_table(rows); continue

    if stripped.startswith('#'):
        lvl = len(stripped) - len(stripped.lstrip('#'))
        txt = stripped[lvl:].strip()
        if lvl == 1 and not first_heading:
            pass
        first_heading = False
        p = doc.add_paragraph(style='Heading %d' % min(lvl, 4))
        add_runs(p, txt)
        i += 1; continue

    if stripped.startswith('---'):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '9AA7B5')
        pbdr.append(bottom); pPr.append(pbdr)
        i += 1; continue

    if stripped.startswith('> '):
        block = []
        while i < len(raw) and raw[i].strip().startswith('>'):
            block.append(raw[i].strip().lstrip('>').strip()); i += 1
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.5); pf.space_before = Pt(5); pf.space_after = Pt(7)
        add_runs(p, ' '.join(x for x in block if x))
        for r in p.runs:
            r.font.size = Pt(10)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr'); left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '6'); left.set(qn('w:color'), '003A70')
        pbdr.append(left); pPr.append(pbdr)
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F2F6FA')
        pPr.append(shd)
        continue

    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
    if m:
        indent = len(m.group(1))
        style = 'List Number' if m.group(2)[0].isdigit() else 'List Bullet'
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = Cm(0.7 + 0.5 * (indent // 2))
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, m.group(3))
        i += 1; continue

    if stripped.startswith('$$'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(stripped.strip('$ '))
        r.italic = True; r.font.size = Pt(10)
        i += 1; continue

    if not stripped:
        i += 1; continue

    p = doc.add_paragraph()
    add_runs(p, stripped)
    i += 1

doc.save(OUT)
print('OK ->', OUT)
