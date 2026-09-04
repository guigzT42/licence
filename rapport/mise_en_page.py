# -*- coding: utf-8 -*-
"""
Bibliothèque de mise en page du rapport (python-docx).
Styles, tableaux, encadrés, figures, en-têtes et pieds de page.
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(BASE, "images")
FIG = os.path.join(BASE, "figures")

# --- Palette (identique au diaporama de soutenance) -----------------------
VERT = RGBColor(0x2E, 0x7D, 0x32)
VERT_CLAIR = RGBColor(0x66, 0xBB, 0x6A)
ORANGE = RGBColor(0xEF, 0x6C, 0x00)
BLEU = RGBColor(0x15, 0x65, 0xC0)
ROUGE = RGBColor(0xC6, 0x28, 0x28)
GRIS = RGBColor(0x54, 0x6E, 0x7A)
ANTHRA = RGBColor(0x37, 0x47, 0x4F)
NOIR = RGBColor(0x21, 0x21, 0x21)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

H_VERT = "2E7D32"
H_VERT_PALE = "EDF5EE"
H_VERT_TRES_PALE = "F5FAF6"
H_ORANGE_PALE = "FFF4E5"
H_BLEU_PALE = "E8F1FB"
H_ROUGE_PALE = "FDECEC"
H_GRIS_PALE = "F2F4F5"
H_GRIS_BORD = "CFD8DC"

POLICE = "Calibri"


# ==========================================================================
#  Utilitaires XML bas niveau
# ==========================================================================
def _shade(element, couleur_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    element.append(shd)


def shade_cell(cell, couleur_hex):
    _shade(cell._tc.get_or_add_tcPr(), couleur_hex)


def shade_paragraph(par, couleur_hex):
    _shade(par._p.get_or_add_pPr(), couleur_hex)


def set_borders(element, cotes=("top", "left", "bottom", "right"),
                sz=6, couleur=H_GRIS_BORD, val="single", space=0):
    pr = element
    borders = OxmlElement("w:pBdr") if pr.tag.endswith("pPr") else OxmlElement("w:tcBorders")
    for cote in cotes:
        e = OxmlElement("w:" + cote)
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), couleur)
        borders.append(e)
    pr.append(borders)


def cell_borders(cell, cotes=("top", "left", "bottom", "right"), sz=6,
                 couleur=H_GRIS_BORD, val="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for cote in cotes:
        e = OxmlElement("w:" + cote)
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), couleur)
        borders.append(e)
    tcPr.append(borders)


def par_border(par, cotes=("left",), sz=18, couleur=H_VERT, space=8):
    pPr = par._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    borders = OxmlElement("w:pBdr")
    for cote in cotes:
        e = OxmlElement("w:" + cote)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), str(space))
        e.set(qn("w:color"), couleur)
        borders.append(e)
    pPr.append(borders)


def keep_with_next(par, valeur=True):
    pPr = par._p.get_or_add_pPr()
    e = OxmlElement("w:keepNext")
    e.set(qn("w:val"), "1" if valeur else "0")
    pPr.append(e)


def no_split_row(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    trPr.append(e)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def add_field(par, instruction):
    """Insère un champ Word (numéro de page, table des matières, ...)."""
    r1 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instruction
    r2._r.append(it)
    r3 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc)
    r4 = par.add_run("1")
    r5 = par.add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r5._r.append(fc)
    return [r1, r2, r3, r4, r5]


def update_fields_on_open(doc):
    """Demande à Word de recalculer les champs (table des matières) à
    l'ouverture. L'élément doit respecter l'ordre du schéma CT_Settings :
    il se place avant <w:compat>, <w:footnotePr> ou <w:hdrShapeDefaults>."""
    settings = doc.settings.element
    for existant in settings.findall(qn("w:updateFields")):
        settings.remove(existant)
    e = OxmlElement("w:updateFields")
    e.set(qn("w:val"), "true")
    for nom in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr",
                "w:compat", "w:docVars", "w:rsids", "w:mathPr"):
        ancre = settings.find(qn(nom))
        if ancre is not None:
            ancre.addprevious(e)
            return
    settings.append(e)


# ==========================================================================
#  Construction du document
# ==========================================================================
class Rapport:

    def __init__(self, titre_courant):
        self.doc = Document()
        self.titre_courant = titre_courant
        self._n1 = 0
        self._n2 = 0
        self._n3 = 0
        self._nfig = 0
        self._ntab = 0
        self.titres = []
        self.ancre_toc = None
        self._setup_page()
        self._setup_styles()

    # ---- mise en page générale -------------------------------------------
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)
        s.top_margin = Cm(1.7)
        s.bottom_margin = Cm(1.5)
        s.header_distance = Cm(1.0)
        s.footer_distance = Cm(0.9)

    def _setup_styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = POLICE
        st.font.size = Pt(10)
        st.font.color.rgb = NOIR
        st.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
        pf = st.paragraph_format
        pf.space_after = Pt(4)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.04
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    @property
    def largeur_utile(self):
        s = self.doc.sections[0]
        return s.page_width - s.left_margin - s.right_margin

    # ---- en-tête / pied de page ------------------------------------------
    def entete_pied(self):
        s = self.doc.sections[0]
        s.different_first_page_header_footer = True

        p = s.header.paragraphs[0]
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(self.titre_courant)
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        r.font.name = POLICE
        par_border(p, cotes=("bottom",), sz=6, couleur=H_GRIS_BORD, space=2)

        f = s.footer.paragraphs[0]
        f.text = ""
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = f.add_run("Guillaume Tardy — Bloc 1 « Réalisation d'études techniques "
                      "pour des bâtiments performants »        Page ")
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        r.font.name = POLICE
        for rr in add_field(f, " PAGE "):
            rr.font.size = Pt(8)
            rr.font.color.rgb = GRIS
            rr.font.name = POLICE
        r = f.add_run(" / ")
        r.font.size = Pt(8)
        r.font.color.rgb = GRIS
        r.font.name = POLICE
        for rr in add_field(f, " NUMPAGES "):
            rr.font.size = Pt(8)
            rr.font.color.rgb = GRIS
            rr.font.name = POLICE

    # ---- primitives de texte ---------------------------------------------
    def _p(self, style=None):
        return self.doc.add_paragraph(style=style)

    def para(self, contenu, size=10, couleur=NOIR, align="justify",
             space_after=4, space_before=0, italique=False, gras=False,
             interligne=1.04, indent=None):
        """contenu : str, ou liste de (texte, gras, couleur[, italique])."""
        p = self._p()
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(space_before)
        pf.line_spacing = interligne
        pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        if indent:
            pf.left_indent = Cm(indent)
        if isinstance(contenu, str):
            contenu = [(contenu, gras, couleur, italique)]
        for item in contenu:
            txt = item[0]
            g = item[1] if len(item) > 1 else False
            c = item[2] if len(item) > 2 else couleur
            it = item[3] if len(item) > 3 else italique
            r = p.add_run(txt)
            r.font.name = POLICE
            r.font.size = Pt(size)
            r.font.bold = g
            r.font.italic = it
            r.font.color.rgb = c
        return p

    # ---- titres ----------------------------------------------------------
    def titre1(self, texte, saut_page=True, numerote=True):
        if saut_page:
            self.doc.add_page_break()
        self._n1 += 1
        self._n2 = 0
        self._n3 = 0
        p = self._p(style="Heading 1")
        pf = p.paragraph_format
        pf.space_before = Pt(0 if saut_page else 13)
        pf.space_after = Pt(7)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        if numerote:
            r = p.add_run(f"{self._n1}.  ")
            r.font.name = POLICE
            r.font.size = Pt(17)
            r.font.bold = True
            r.font.color.rgb = VERT_CLAIR
        self.titres.append((1, (f"{self._n1}.  " if numerote else "") + texte))
        r = p.add_run(texte.upper())
        r.font.name = POLICE
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = VERT
        par_border(p, cotes=("bottom",), sz=12, couleur=H_VERT, space=5)
        return p

    def titre2(self, texte):
        self._n2 += 1
        self._n3 = 0
        p = self._p(style="Heading 2")
        pf = p.paragraph_format
        pf.space_before = Pt(11)
        pf.space_after = Pt(4)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        self.titres.append((2, f"{self._n1}.{self._n2}   {texte}"))
        r = p.add_run(f"{self._n1}.{self._n2}   {texte}")
        r.font.name = POLICE
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = VERT
        return p

    def titre3(self, texte):
        self._n3 += 1
        p = self._p(style="Heading 3")
        pf = p.paragraph_format
        pf.space_before = Pt(9)
        pf.space_after = Pt(3)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        r = p.add_run(f"{self._n1}.{self._n2}.{self._n3}   {texte}")
        r.font.name = POLICE
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = ANTHRA
        return p

    def sous_titre(self, texte, couleur=ORANGE):
        p = self._p()
        pf = p.paragraph_format
        pf.space_before = Pt(9)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(texte)
        r.font.name = POLICE
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = couleur
        return p

    # ---- listes ----------------------------------------------------------
    def puces(self, items, size=10, couleur=NOIR, puce="▪",
              couleur_puce=VERT, indent=0.5, space_after=2):
        """items : str, ou (label_gras, suite)."""
        for it in items:
            p = self._p()
            pf = p.paragraph_format
            pf.left_indent = Cm(indent + 0.45)
            pf.first_line_indent = Cm(-0.45)
            pf.space_after = Pt(space_after)
            pf.line_spacing = 1.04
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(puce + "   ")
            r.font.name = POLICE
            r.font.size = Pt(size)
            r.font.color.rgb = couleur_puce
            r.font.bold = True
            if isinstance(it, str):
                it = [(it, False)]
            elif isinstance(it, tuple):
                it = [(it[0], True), (it[1], False)]
            for seg in it:
                r = p.add_run(seg[0])
                r.font.name = POLICE
                r.font.size = Pt(size)
                r.font.bold = seg[1] if len(seg) > 1 else False
                r.font.color.rgb = seg[2] if len(seg) > 2 else couleur

    def numerotees(self, items, size=10.5, indent=0.5):
        for i, it in enumerate(items, 1):
            p = self._p()
            pf = p.paragraph_format
            pf.left_indent = Cm(indent + 0.55)
            pf.first_line_indent = Cm(-0.55)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.04
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(f"{i}.   ")
            r.font.name = POLICE
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = VERT
            if isinstance(it, str):
                it = [(it, False)]
            elif isinstance(it, tuple):
                it = [(it[0], True), (it[1], False)]
            for seg in it:
                r = p.add_run(seg[0])
                r.font.name = POLICE
                r.font.size = Pt(size)
                r.font.bold = seg[1] if len(seg) > 1 else False

    # ---- encadrés --------------------------------------------------------
    def encadre(self, titre, lignes, couleur=VERT, fond=H_VERT_PALE,
                size=9, icone=None):
        """Encadré d'une cellule, avec liseré de couleur à gauche."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell = t.cell(0, 0)
        cell.width = self.largeur_utile
        shade_cell(cell, fond)
        cell_borders(cell, cotes=("top", "bottom", "right"), sz=4,
                     couleur=H_GRIS_BORD)
        cell_borders_left = OxmlElement("w:tcBorders")
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        for cote, sz, col in (("top", 4, H_GRIS_BORD), ("bottom", 4, H_GRIS_BORD),
                              ("right", 4, H_GRIS_BORD),
                              ("left", 30, "%02X%02X%02X" % (couleur[0] if isinstance(couleur, tuple) else 0, 0, 0))):
            pass
        borders = OxmlElement("w:tcBorders")
        hexcol = "{:02X}{:02X}{:02X}".format(*_rgb_tuple(couleur))
        for cote, sz, col in (("top", 4, H_GRIS_BORD), ("bottom", 4, H_GRIS_BORD),
                              ("right", 4, H_GRIS_BORD), ("left", 28, hexcol)):
            e = OxmlElement("w:" + cote)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), col)
            borders.append(e)
        tcPr.append(borders)
        self._marges_cellule(cell, haut=80, bas=80, gauche=150, droite=120)

        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(3 if lignes else 0)
        p0.paragraph_format.line_spacing = 1.1
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if titre:
            r = p0.add_run((icone + "  " if icone else "") + titre)
            r.font.name = POLICE
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = couleur
        for i, ligne in enumerate(lignes):
            p = cell.add_paragraph() if (titre or i) else p0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.04
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if isinstance(ligne, str):
                ligne = [(ligne, False)]
            elif isinstance(ligne, tuple):
                ligne = [(ligne[0], True), (ligne[1], False)]
            for seg in ligne:
                r = p.add_run(seg[0])
                r.font.name = POLICE
                r.font.size = Pt(size)
                r.font.bold = seg[1] if len(seg) > 1 else False
                r.font.color.rgb = seg[2] if len(seg) > 2 else NOIR
        self.espace(4)
        return t

    def _marges_cellule(self, cell, haut=80, bas=80, gauche=120, droite=120):
        tcPr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for nom, val in (("top", haut), ("bottom", bas),
                         ("start", gauche), ("end", droite)):
            e = OxmlElement("w:" + nom)
            e.set(qn("w:w"), str(val))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tcPr.append(mar)

    def kpi(self, valeurs):
        """valeurs : liste de (valeur, libellé, couleur)."""
        n = len(valeurs)
        t = self.doc.add_table(rows=1, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        larg = int(self.largeur_utile / n)
        for i, (val, lib, coul) in enumerate(valeurs):
            c = t.cell(0, i)
            c.width = Emu(larg)
            shade_cell(c, H_GRIS_PALE)
            hexcol = "{:02X}{:02X}{:02X}".format(*_rgb_tuple(coul))
            tcPr = c._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn("w:tcBorders")):
                tcPr.remove(old)
            borders = OxmlElement("w:tcBorders")
            for cote, sz, col in (("top", 26, hexcol), ("bottom", 4, "FFFFFF"),
                                  ("left", 12, "FFFFFF"), ("right", 12, "FFFFFF")):
                e = OxmlElement("w:" + cote)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), col)
                borders.append(e)
            tcPr.append(borders)
            self._marges_cellule(c, haut=80, bas=80, gauche=70, droite=70)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            r.font.name = POLICE
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = coul
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r = p2.add_run(lib)
            r.font.name = POLICE
            r.font.size = Pt(8.5)
            r.font.color.rgb = GRIS
        self.espace(5)
        return t

    # ---- tableaux --------------------------------------------------------
    def tableau(self, lignes, largeurs=None, size=8, size_entete=8,
                couleur_entete=H_VERT, premiere_col_gras=True,
                lignes_surlignees=(), fond_surlignage=H_VERT_PALE,
                titre=None, align_centre_cols=(), zebra=True):
        if titre:
            self._ntab += 1
            p = self.para([(f"Tableau {self._ntab} — ", True, VERT),
                           (titre, False, ANTHRA)],
                          size=9, align="left", space_after=3, space_before=6)
            keep_with_next(p)

        nl, nc = len(lignes), len(lignes[0])
        t = self.doc.add_table(rows=nl, cols=nc)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if largeurs is None:
            largeurs = [1.0 / nc] * nc
        tot = self.largeur_utile
        larg_emu = [Emu(int(tot * w)) for w in largeurs]

        for i, ligne in enumerate(lignes):
            row = t.rows[i]
            if i == 0:
                no_split_row(row)
                repeat_header(row)
            for j, val in enumerate(ligne):
                c = t.cell(i, j)
                c.width = larg_emu[j]
                self._marges_cellule(c, haut=28, bas=28, gauche=70, droite=70)
                if i == 0:
                    shade_cell(c, couleur_entete)
                elif i in lignes_surlignees:
                    shade_cell(c, fond_surlignage)
                elif zebra and i % 2 == 0:
                    shade_cell(c, H_GRIS_PALE)
                cell_borders(c, sz=4, couleur="FFFFFF" if i == 0 else H_GRIS_BORD)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if j in align_centre_cols or (j > 0 and largeurs[j] < 0.13):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                segments = val if isinstance(val, list) else [(str(val), None)]
                for seg in segments:
                    r = p.add_run(seg[0])
                    r.font.name = POLICE
                    r.font.size = Pt(size_entete if i == 0 else size)
                    if i == 0:
                        r.font.bold = True
                        r.font.color.rgb = BLANC
                    else:
                        r.font.bold = (seg[1] if len(seg) > 1 and seg[1] is not None
                                       else (premiere_col_gras and j == 0)
                                       or i in lignes_surlignees)
                        r.font.color.rgb = (seg[2] if len(seg) > 2 else
                                            (VERT if i in lignes_surlignees else NOIR))
        self.espace(6)
        return t

    # ---- figures ---------------------------------------------------------
    def figure(self, chemin, largeur_cm=14.0, legende=None, centrer=True,
               numerote=True):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrer else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        keep_with_next(p)
        p.add_run().add_picture(chemin, width=Cm(largeur_cm))
        if legende:
            if numerote:
                self._nfig += 1
                txt = [(f"Figure {self._nfig} — ", True, VERT),
                       (legende, False, GRIS, True)]
            else:
                txt = [(legende, False, GRIS, True)]
            lp = self.para(txt, size=8.5, align="center", space_after=8)
        return p

    def images_cote_a_cote(self, images, hauteur_cm=None, largeurs_cm=None,
                           legende_globale=None):
        """images : liste de (chemin, légende)."""
        n = len(images)
        t = self.doc.add_table(rows=2, cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        larg = int(self.largeur_utile / n)
        for i, (chemin, leg) in enumerate(images):
            c = t.cell(0, i)
            c.width = Emu(larg)
            self._marges_cellule(c, haut=25, bas=20, gauche=50, droite=50)
            cell_borders(c, sz=0, couleur="FFFFFF", val="none")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run()
            if largeurs_cm:
                run.add_picture(chemin, width=Cm(largeurs_cm[i]))
            elif hauteur_cm:
                run.add_picture(chemin, height=Cm(hauteur_cm))
            else:
                run.add_picture(chemin, width=Cm(self.largeur_utile.cm / n - 0.4))
            c2 = t.cell(1, i)
            c2.width = Emu(larg)
            self._marges_cellule(c2, haut=6, bas=25, gauche=50, droite=50)
            cell_borders(c2, sz=0, couleur="FFFFFF", val="none")
            p2 = c2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(0)
            r = p2.add_run(leg)
            r.font.name = POLICE
            r.font.size = Pt(8.5)
            r.font.italic = True
            r.font.color.rgb = GRIS
        if legende_globale:
            self._nfig += 1
            self.para([(f"Figure {self._nfig} — ", True, VERT),
                       (legende_globale, False, GRIS, True)],
                      size=8.5, align="center", space_after=8)
        else:
            self.espace(5)
        return t

    # ---- divers ----------------------------------------------------------
    def espace(self, pts=8):
        p = self._p()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run("")
        r.font.size = Pt(pts / 2.0)
        return p

    def saut_page(self):
        self.doc.add_page_break()

    def calcul(self, lignes, titre=None, couleur=BLEU):
        """Bloc de calcul encadré, police à chasse fixe."""
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        cell = t.cell(0, 0)
        cell.width = self.largeur_utile
        shade_cell(cell, "F7F9FA")
        hexcol = "{:02X}{:02X}{:02X}".format(*_rgb_tuple(couleur))
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        borders = OxmlElement("w:tcBorders")
        for cote, sz, col in (("top", 4, H_GRIS_BORD), ("bottom", 4, H_GRIS_BORD),
                              ("right", 4, H_GRIS_BORD), ("left", 24, hexcol)):
            e = OxmlElement("w:" + cote)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), col)
            borders.append(e)
        tcPr.append(borders)
        self._marges_cellule(cell, haut=70, bas=70, gauche=150, droite=120)
        first = True
        if titre:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(titre)
            r.font.name = POLICE
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = couleur
            first = False
        for ligne in lignes:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            gras = ligne.startswith("=>")
            txt = ligne[2:].strip() if gras else ligne
            r = p.add_run(("→  " if gras else "") + txt)
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            r.font.bold = gras
            r.font.color.rgb = VERT if gras else ANTHRA
        self.espace(5)
        return t

    def toc_ancre(self):
        """Réserve l'emplacement de la table des matières."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        self.ancre_toc = p
        return p

    def remplir_toc(self):
        """Insère la table des matières à l'emplacement réservé.

        Le contenu statique est encapsulé dans un champ TOC : il s'affiche tel
        quel dans tout lecteur, et Word le recalcule avec les numéros de page
        à l'ouverture du document.
        """
        if self.ancre_toc is None or not self.titres:
            return
        ancre = self.ancre_toc._p
        pars = []
        for i, (niveau, texte) in enumerate(self.titres):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(2 if niveau == 1 else 1)
            pf.space_before = Pt(6 if niveau == 1 and i else 0)
            pf.left_indent = Cm(0 if niveau == 1 else 0.7)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if i == 0:
                r1 = p.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "begin")
                r1._r.append(fc)
                r2 = p.add_run()
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = r' TOC \o "1-2" \h \z \u '
                r2._r.append(it)
                r3 = p.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "separate")
                r3._r.append(fc)
            r = p.add_run(texte)
            r.font.name = POLICE
            r.font.size = Pt(11 if niveau == 1 else 9.5)
            r.font.bold = (niveau == 1)
            r.font.color.rgb = VERT if niveau == 1 else GRIS
            if i == len(self.titres) - 1:
                r4 = p.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "end")
                r4._r.append(fc)
            pars.append(p)
        for p in pars:
            ancre.addprevious(p._p)
        ancre.getparent().remove(ancre)

    def enregistrer(self, chemin):
        update_fields_on_open(self.doc)
        self.doc.save(chemin)
        return chemin


def _rgb_tuple(c):
    if isinstance(c, RGBColor):
        s = str(c)
        return (int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
    return c
