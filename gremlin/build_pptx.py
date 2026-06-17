# -*- coding: utf-8 -*-
"""
Diaporama de soutenance - Projet GREMLIN (phase de préparation de la maquette).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

IMG = "/home/user/licence/gremlin/img"
OUT = "/home/user/licence/Projet_GREMLIN_Presentation.pptx"

# ---------- Charte ----------
NAVY   = RGBColor(0x1F, 0x4E, 0x79)
NAVY2  = RGBColor(0x2E, 0x6B, 0xA8)
LBLUE  = RGBColor(0xEA, 0xF1, 0xF8)
GREY   = RGBColor(0x40, 0x40, 0x40)
LGREY  = RGBColor(0x8A, 0x8A, 0x8A)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xC0, 0x39, 0x2B)
GREEN  = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
BG     = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

NBSP = " "
def fr(t):
    for a, b in ((" :", NBSP+":"), (" ;", NBSP+";"), (" !", NBSP+"!"),
                 (" ?", NBSP+"?"), (" %", NBSP+"%"), ("« ", "«"+NBSP), (" »", NBSP+"»")):
        t = t.replace(a, b)
    return t

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    return s

def rect(s, l, t, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def rrect(s, l, t, w, h, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1.25)
    sp.shadow.inherit = False
    return sp

def txt(s, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: liste de dicts {text,size,bold,color,italic,space_after,bullet}."""
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        p.space_before = Pt(ln.get("space_before", 0))
        if ln.get("level"):
            p.level = ln["level"]
        r = p.add_run(); r.text = fr(ln["text"])
        f = r.font
        f.size = Pt(ln.get("size", 18)); f.bold = ln.get("bold", False)
        f.italic = ln.get("italic", False); f.name = "Calibri"
        f.color.rgb = ln.get("color", GREY)
    return tb

def title_bar(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.25), NAVY)
    rect(s, 0, Inches(1.25), SW, Pt(3), ORANGE)
    txt(s, Inches(0.55), Inches(0.16), Inches(11.5), Inches(0.35),
        [{"text": kicker.upper(), "size": 12, "bold": True, "color": RGBColor(0xBE,0xD3,0xE9)}])
    txt(s, Inches(0.55), Inches(0.45), Inches(12.2), Inches(0.7),
        [{"text": title, "size": 26, "bold": True, "color": WHITE}])

def footer(s, n):
    txt(s, Inches(0.55), Inches(7.05), Inches(6), Inches(0.3),
        [{"text": "Projet GREMLIN — Soutenance", "size": 9, "color": LGREY}])
    txt(s, Inches(11.3), Inches(7.05), Inches(1.5), Inches(0.3),
        [{"text": str(n), "size": 9, "color": LGREY, "align": PP_ALIGN.RIGHT}])

def add_pic(s, path, l, t, w=None, h=None):
    im = Image.open(path); iw, ih = im.size; ar = iw/ih
    if w and not h:
        h = Emu(int(w * ih / iw))
    elif h and not w:
        w = Emu(int(h * iw / ih))
    s.shapes.add_picture(path, l, t, width=w, height=h)
    return w, h

def pic_centered(s, path, top, max_w, max_h):
    im = Image.open(path); iw, ih = im.size
    w = max_w; h = Emu(int(w * ih / iw))
    if h > max_h:
        h = max_h; w = Emu(int(h * iw / ih))
    left = Emu(int((SW - w) / 2))
    s.shapes.add_picture(path, left, top, width=w, height=h)
    return w, h

def kpi(s, l, t, w, value, label, color):
    box = rrect(s, l, t, w, Inches(1.35), LBLUE)
    txt(s, l, t+Inches(0.16), w, Inches(0.7),
        [{"text": value, "size": 30, "bold": True, "color": color, "align": PP_ALIGN.CENTER}])
    txt(s, l, t+Inches(0.86), w, Inches(0.45),
        [{"text": label, "size": 12.5, "color": GREY, "align": PP_ALIGN.CENTER}])

def bullets(s, l, t, w, h, items, size=18, gap=8):
    lines = []
    for it in items:
        if isinstance(it, tuple):
            text, lvl = it
        else:
            text, lvl = it, 0
        lines.append({"text": ("•  " if lvl == 0 else "–  ") + text,
                      "size": size if lvl == 0 else size-2,
                      "color": GREY if lvl == 0 else LGREY,
                      "space_after": gap, "level": lvl})
    txt(s, l, t, w, h, lines)

# =====================================================================
# 1. TITRE
# =====================================================================
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(3.05), SW, Pt(3), ORANGE)
rect(s, Inches(0.0), Inches(0.0), Inches(0.22), SH, ORANGE)
txt(s, Inches(1.0), Inches(0.9), Inches(11), Inches(0.5),
    [{"text": "CONDUITE DE PROJET — CAS N° 1", "size": 15, "bold": True,
      "color": RGBColor(0xBE,0xD3,0xE9)}])
txt(s, Inches(1.0), Inches(1.5), Inches(11.5), Inches(1.6),
    [{"text": "Projet GREMLIN", "size": 58, "bold": True, "color": WHITE}])
txt(s, Inches(1.0), Inches(3.25), Inches(11.5), Inches(1.2),
    [{"text": "Essai en soufflerie d’une maquette", "size": 26, "color": RGBColor(0xDD,0xE8,0xF3)},
     {"text": "Phase de préparation de la maquette", "size": 18, "italic": True,
      "color": RGBColor(0xBE,0xD3,0xE9), "space_before": 6}])
txt(s, Inches(1.0), Inches(5.7), Inches(11.5), Inches(1.3),
    [{"text": "Scénario · Organigramme des tâches · Planning · Budget · Risques · Pilotage",
      "size": 14, "color": RGBColor(0x9F,0xBA,0xD6)},
     {"text": "Chargé de projet — 17 juin 2026", "size": 13, "color": RGBColor(0x9F,0xBA,0xD6),
      "space_before": 10}])

# =====================================================================
# 2. CONTEXTE & ENJEU
# =====================================================================
s = slide(); title_bar(s, "Cadrage", "Contexte et enjeu"); footer(s, 2)
bullets(s, Inches(0.6), Inches(1.7), Inches(7.0), Inches(5),
    [ "Un avionneur étranger de notoriété internationale nous confie l’essai en "
      "soufflerie d’une maquette (nouvel avion, projet GREMLIN).",
      ("Opération modeste, mais nouveau client à fort potentiel : la Direction veut "
       "absolument décrocher la commande.", 1),
      "Devis initial : 297 K€. Le client le juge élevé (concurrent à 259 K€)…",
      ("… mais il est sensible au sérieux et à l’engagement sur les délais.", 1),
      "Après négociation, commande conclue.",
    ], size=17, gap=12)
# Encadré commande
rrect(s, Inches(8.0), Inches(1.9), Inches(4.7), Inches(2.0), LBLUE)
txt(s, Inches(8.0), Inches(2.05), Inches(4.7), Inches(0.4),
    [{"text": "LA COMMANDE", "size": 13, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER}])
txt(s, Inches(8.2), Inches(2.55), Inches(4.3), Inches(1.3),
    [{"text": "272 000 € — budget global", "size": 18, "bold": True, "color": GREY, "space_after":8},
     {"text": "85 jours ouvrables — délai ferme", "size": 18, "bold": True, "color": GREY}])
rrect(s, Inches(8.0), Inches(4.1), Inches(4.7), Inches(2.4), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(8.2), Inches(4.3), Inches(4.3), Inches(2.1),
    [{"text": "L’enjeu", "size": 14, "bold": True, "color": RED, "space_after": 6},
     {"text": "Réussir ce premier essai pour ouvrir une relation durable avec un "
      "client capable de générer des débouchés importants.", "size": 14.5, "color": GREY}])

# =====================================================================
# 3. PÉRIMÈTRE & CONTRAINTES
# =====================================================================
s = slide(); title_bar(s, "Cadrage", "Périmètre étudié et contraintes"); footer(s, 3)
txt(s, Inches(0.6), Inches(1.55), Inches(12), Inches(0.8),
    [{"text": "Ce dossier traite la 1re phase du projet : la préparation de la maquette "
      "— la fixer sur son support et l’équiper en capteurs, jusqu’à une maquette "
      "« prête à l’expérimentation ».", "size": 16.5, "color": GREY}])
# 3 KPI contraintes
kpi(s, Inches(0.6), Inches(2.7), Inches(3.9), "≤ 45 j", "Délai de la phase", NAVY)
kpi(s, Inches(4.7), Inches(2.7), Inches(3.9), "≤ 48 500 €", "Budget de la phase", NAVY)
kpi(s, Inches(8.8), Inches(2.7), Inches(3.9), "Qualité", "Essai réussi + DVD", NAVY)
txt(s, Inches(0.6), Inches(4.5), Inches(12), Inches(2),
    [{"text": "Pourquoi cette phase est sensible", "size": 16, "bold": True, "color": NAVY,
      "space_after": 8}])
bullets(s, Inches(0.6), Inches(5.0), Inches(12), Inches(2),
    [ "C’est le point de départ : si elle dérape, tout l’engagement des 85 jours est menacé.",
      "Les pièces de fixation sont fabriquées en interne ; les équipements de mesure sont achetés à l’extérieur.",
      "L’équipe de préparation ne participe pas aux mesures : elle laisse un mode opératoire.",
    ], size=16, gap=8)

# =====================================================================
# 4. SCÉNARIO
# =====================================================================
s = slide(); title_bar(s, "Livrable 1", "Scénario du projet"); footer(s, 4)
txt(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.6),
    [{"text": "Cinq grandes étapes ; trois peuvent avancer en parallèle, ce qui sera "
      "la clé pour tenir le délai.", "size": 15.5, "color": GREY}])
pic_centered(s, IMG + "/scenario.png", Inches(2.15), Inches(12.4), Inches(4.7))

# =====================================================================
# 5. ORGANIGRAMME DES TÂCHES
# =====================================================================
s = slide(); title_bar(s, "Livrable 2", "Organigramme des tâches (WBS)"); footer(s, 5)
txt(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.6),
    [{"text": "On part du produit final (maquette prête) que l’on décompose en 6 lots, "
      "puis en 22 tâches élémentaires. En rouge : les tâches qui seront critiques.",
      "size": 15, "color": GREY}])
pic_centered(s, IMG + "/wbs.png", Inches(2.2), Inches(12.0), Inches(4.7))

# =====================================================================
# 6. PLANNING — PERT / CHEMIN CRITIQUE
# =====================================================================
s = slide(); title_bar(s, "Livrable 3", "Planning — réseau et chemin critique"); footer(s, 6)
pic_centered(s, IMG + "/pert.png", Inches(1.55), Inches(12.6), Inches(4.05))
rrect(s, Inches(0.7), Inches(5.85), Inches(11.95), Inches(1.05), LBLUE)
txt(s, Inches(0.9), Inches(6.02), Inches(11.6), Inches(0.8),
    [{"text": "Chemin critique : A → G → H → I → J → S → T → U → V", "size": 16, "bold": True,
      "color": NAVY, "space_after": 2},
     {"text": "Tâche la plus dangereuse : H « fabrication des pièces » (14 jours, aucune marge).",
      "size": 13.5, "color": GREY}])

# =====================================================================
# 7. PLANNING — GANTT + VERDICT
# =====================================================================
s = slide(); title_bar(s, "Livrable 3", "Planning — Gantt et vérification du délai"); footer(s, 7)
pic_centered(s, IMG + "/gantt.png", Inches(1.5), Inches(8.7), Inches(4.7))
# Verdict à droite
rrect(s, Inches(9.55), Inches(1.6), Inches(3.25), Inches(2.2), RGBColor(0xE8,0xF5,0xE9))
txt(s, Inches(9.7), Inches(1.75), Inches(2.95), Inches(2),
    [{"text": "DÉLAI", "size": 13, "bold": True, "color": GREEN, "align": PP_ALIGN.CENTER, "space_after":4},
     {"text": "45 / 45 j", "size": 30, "bold": True, "color": GREEN, "align": PP_ALIGN.CENTER, "space_after":4},
     {"text": "Contrainte respectée", "size": 13, "color": GREY, "align": PP_ALIGN.CENTER}])
rrect(s, Inches(9.55), Inches(4.0), Inches(3.25), Inches(2.6), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(9.7), Inches(4.18), Inches(2.95), Inches(2.3),
    [{"text": "Le point d’alerte", "size": 14, "bold": True, "color": RED, "align": PP_ALIGN.CENTER, "space_after":6},
     {"text": "Zéro marge sur le chemin critique. Le moindre retard repousse la fin de phase.",
      "size": 13.5, "color": GREY, "align": PP_ALIGN.CENTER}])

# =====================================================================
# 8. BUDGET — SYNTHÈSE + VERDICT
# =====================================================================
s = slide(); title_bar(s, "Livrable 4", "Budget — synthèse"); footer(s, 8)
# Tableau synthèse à gauche
rows = [
    ("Main-d’œuvre (Ing. + Tech. + Ach. + Ouv.)", "41 796 €"),
    ("Fabrication des pièces (atelier)", "2 940 €"),
    ("Équipements achetés à l’extérieur", "5 200 €"),
    ("COÛT TOTAL DE LA PHASE", "49 936 €"),
    ("Enveloppe allouée", "48 500 €"),
]
y = Inches(1.75)
for i, (lab, val) in enumerate(rows):
    is_tot = lab.startswith("COÛT")
    bg = RGBColor(0xFD,0xEC,0xEA) if is_tot else (LBLUE if i % 2 == 0 else WHITE)
    rect(s, Inches(0.6), y, Inches(7.4), Inches(0.62), bg)
    txt(s, Inches(0.8), y+Inches(0.10), Inches(5.4), Inches(0.5),
        [{"text": lab, "size": 14, "bold": is_tot, "color": NAVY if is_tot else GREY}])
    txt(s, Inches(6.0), y+Inches(0.10), Inches(1.85), Inches(0.5),
        [{"text": val, "size": 14, "bold": is_tot, "color": RED if is_tot else GREY,
          "align": PP_ALIGN.RIGHT}])
    y = y + Inches(0.62)
# Verdict
rrect(s, Inches(8.3), Inches(1.75), Inches(4.45), Inches(2.4), RGBColor(0xFD,0xEC,0xEA))
txt(s, Inches(8.5), Inches(1.95), Inches(4.05), Inches(2.1),
    [{"text": "ÉCART", "size": 14, "bold": True, "color": RED, "align": PP_ALIGN.CENTER, "space_after":4},
     {"text": "+ 1 436 €", "size": 34, "bold": True, "color": RED, "align": PP_ALIGN.CENTER, "space_after":4},
     {"text": "Dépassement ≈ 3 % : la phase est légèrement déficitaire.",
      "size": 13.5, "color": GREY, "align": PP_ALIGN.CENTER}])
txt(s, Inches(8.3), Inches(4.45), Inches(4.45), Inches(2.3),
    [{"text": "Pistes de retour à l’équilibre", "size": 14.5, "bold": True, "color": NAVY, "space_after":6}])
bullets(s, Inches(8.3), Inches(4.95), Inches(4.5), Inches(2),
    [ "Alléger les grosses tâches ouvrier (S, F, J).",
      "Renégocier équipements / devis atelier.",
      "Basculer des heures ingénieur vers technicien.",
    ], size=13, gap=6)

# =====================================================================
# 9. BUDGET — COURBE EN S
# =====================================================================
s = slide(); title_bar(s, "Livrable 4", "Budget — courbe en « S »"); footer(s, 9)
txt(s, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.6),
    [{"text": "Dépenses cumulées au fil du planning. La courbe franchit l’enveloppe en "
      "toute fin de phase : c’est le dépassement de 1 436 €.", "size": 15.5, "color": GREY}])
pic_centered(s, IMG + "/courbe_s.png", Inches(2.15), Inches(10.5), Inches(4.7))

# =====================================================================
# 10. RISQUES
# =====================================================================
s = slide(); title_bar(s, "Livrable 5", "Analyse des risques"); footer(s, 10)
# Matrice à gauche (placement explicite, calé en hauteur)
_im = Image.open(IMG + "/matrice_risques.png"); _iw, _ih = _im.size
_h = Inches(5.0); _w = Emu(int(_h * _iw / _ih))
s.shapes.add_picture(IMG + "/matrice_risques.png", Inches(0.5), Inches(1.7), width=_w, height=_h)
txt(s, Inches(8.2), Inches(1.7), Inches(4.6), Inches(0.5),
    [{"text": "3 risques prioritaires (criticité ≥ 9)", "size": 16, "bold": True, "color": NAVY}])
prio = [
    ("R1", "Retard de fabrication des pièces (tâche H)", "12", RED),
    ("R2", "Modifs logiciels dans un délai serré", "9", ORANGE),
    ("R7", "Aucune marge sur le chemin critique", "9", ORANGE),
]
y = Inches(2.35)
for code, lab, c, col in prio:
    rrect(s, Inches(8.2), y, Inches(4.6), Inches(1.05), LBLUE)
    rect(s, Inches(8.2), y, Inches(0.16), Inches(1.05), col)
    txt(s, Inches(8.45), y+Inches(0.1), Inches(3.2), Inches(0.9),
        [{"text": code, "size": 14, "bold": True, "color": col, "space_after": 2},
         {"text": lab, "size": 12.5, "color": GREY}])
    txt(s, Inches(11.7), y+Inches(0.22), Inches(0.95), Inches(0.6),
        [{"text": c, "size": 22, "bold": True, "color": col, "align": PP_ALIGN.CENTER}])
    y = y + Inches(1.2)

# =====================================================================
# 11. PILOTAGE
# =====================================================================
s = slide(); title_bar(s, "Livrable 6", "Pilotage de la réalisation"); footer(s, 11)
txt(s, Inches(0.6), Inches(1.6), Inches(6), Inches(0.5),
    [{"text": "Tenir le cap : suivi simple et réactif", "size": 17, "bold": True, "color": NAVY}])
bullets(s, Inches(0.6), Inches(2.2), Inches(6.4), Inches(4),
    [ "Revues régulières avec le client, calées sur les jalons.",
      "Mise à jour hebdomadaire du Gantt.",
      "Surveillance prioritaire des tâches critiques (marge nulle).",
      "Utiliser les marges « support » (8 j) et « appro » (3 j) comme amortisseur.",
      "Suivi des dépenses vs courbe en S.",
    ], size=15.5, gap=11)
# Jalons à droite
txt(s, Inches(7.4), Inches(1.6), Inches(5.4), Inches(0.5),
    [{"text": "Jalons clés", "size": 17, "bold": True, "color": NAVY}])
jal = [("J1","Étude validée","j10"),("J2","Pièces prêtes","j30"),
       ("J3","Maquette fixée","j34"),("J4","Revue client","j39"),
       ("J5","Maquette prête","j45")]
y = Inches(2.2)
for code, lab, when in jal:
    rrect(s, Inches(7.4), y, Inches(5.4), Inches(0.66), LBLUE)
    txt(s, Inches(7.6), y+Inches(0.12), Inches(0.7), Inches(0.5),
        [{"text": code, "size": 14, "bold": True, "color": NAVY2}])
    txt(s, Inches(8.3), y+Inches(0.12), Inches(3.2), Inches(0.5),
        [{"text": lab, "size": 14, "color": GREY}])
    txt(s, Inches(11.6), y+Inches(0.12), Inches(1.0), Inches(0.5),
        [{"text": when, "size": 14, "bold": True, "color": NAVY, "align": PP_ALIGN.RIGHT}])
    y = y + Inches(0.78)

# =====================================================================
# 12. CONCLUSION
# =====================================================================
s = slide(); title_bar(s, "Synthèse", "Conclusion"); footer(s, 12)
kpi(s, Inches(0.6), Inches(1.7), Inches(3.9), "45 / 45 j", "Délai : respecté, sans marge", GREEN)
kpi(s, Inches(4.7), Inches(1.7), Inches(3.9), "+ 1 436 €", "Budget : à corriger", RED)
kpi(s, Inches(8.8), Inches(1.7), Inches(3.9), "3", "Risques prioritaires", ORANGE)
txt(s, Inches(0.6), Inches(3.5), Inches(12), Inches(0.5),
    [{"text": "Le projet est tenable, mais ne laisse aucune place à l’improvisation.",
      "size": 18, "bold": True, "color": NAVY}])
txt(s, Inches(0.6), Inches(4.15), Inches(12), Inches(0.5),
    [{"text": "Trois conditions pour honorer l’engagement des 85 jours :", "size": 16, "color": GREY}])
bullets(s, Inches(0.9), Inches(4.75), Inches(11.5), Inches(2),
    [ "sécuriser la fabrication des pièces (tâche H, point dur du planning) ;",
      "faire valider le plan d’économies par la Direction avant le lancement ;",
      "tenir les revues avec le client pour verrouiller les jalons.",
    ], size=16, gap=10)
rect(s, Inches(0.6), Inches(6.75), SW-Inches(1.2), Pt(2.5), ORANGE)
txt(s, Inches(0.6), Inches(6.55), Inches(12), Inches(0.4),
    [{"text": "Au-delà de l’essai : transformer ce premier contrat en relation durable.",
      "size": 13.5, "italic": True, "color": LGREY}])

prs.save(OUT)
print("Présentation générée :", OUT, "—", len(prs.slides._sldIdLst), "diapos")
